"""Build the short unpruuf Compliance sales deck (6 slides, 16:9) in EN and RO, PPTX + PDF.

Usage: python3 build_deck.py            (writes next to this script)
Needs python-pptx, Pillow, LibreOffice (soffice) for the PDFs, fonts Open Sans and DejaVu Sans Mono.
Rules: unpruuf/vertrieb/RULES_FOR_MATERIAL.md, design: design/README.md (NexonAI design, light only).
No prices on the slides. Mailbox picture without numbers. No technology after slide 2.
"""
from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageChops
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
LOGO_SRC = REPO / "design" / "assets" / "nexonai-logo.webp"

# NexonAI tokens (design/nexonai.css)
PRIMARY = "307098"
NAVY = "0f2a3a"
SKY = "90c8e8"
TINT = "e3eef6"
WARM = "a8500f"
WARM_BG = "f8e7d6"
SURFACE = "f7f9fa"
SURFACE_LOW = "f2f4f6"
SURFACE_MID = "eaedef"
WHITE = "ffffff"
INK = "172733"
VARIANT = "505b63"
LINE = "d9dfe3"
RED, AMBER, GREEN = "b3261e", "b7791f", "2e7d4f"
SANS = "Open Sans"
MONO = "DejaVu Sans Mono"

W, H = 13.333, 7.5
MX = 0.6  # side margin


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.upper())


# ---------------------------------------------------------------- primitives
def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=None, shape=None, line_w=0.75):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def card(slide, x, y, w, h, fill=WHITE, line=LINE):
    return rect(slide, x, y, w, h, fill=fill, line=line, radius=0.167)  # 16 px radius


def text(slide, x, y, w, h, paras, size=14, color=INK, font=SANS, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.15, space_after=0):
    """paras: str or list of str / list of (runs) where runs = [(text, {opts})]."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        par.line_spacing = spacing
        par.space_after = Pt(space_after)
        runs = p if isinstance(p, list) else [(p, {})]
        for t, o in runs:
            r = par.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def label(slide, x, y, w, t, color=PRIMARY, size=10.5, align=PP_ALIGN.LEFT):
    return text(slide, x, y, w, 0.3, t.upper(), size=size, color=color, font=MONO, bold=True, align=align)


def chip(slide, x, y, t, fg=PRIMARY, bg=TINT, size=10, pad=0.16, h=0.34):
    w = len(t) * size * 0.0085 + 2 * pad  # mono width estimate
    rect(slide, x, y, w, h, fill=bg, radius=h / 2)
    text(slide, x, y, w, h, t, size=size, color=fg, font=MONO, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    return w


def arrow(slide, x1, y1, x2, y2, color=PRIMARY, width=2.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ---------------------------------------------------------------- frame
def logo_png() -> str:
    im = Image.open(LOGO_SRC).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    diff = ImageChops.difference(im, bg).convert("L").point(lambda v: 255 if v > 18 else 0)
    box = diff.getbbox()
    pad = 6
    im = im.crop((max(0, box[0] - pad), max(0, box[1] - pad), min(im.width, box[2] + pad), min(im.height, box[3] + pad)))
    out = Path(tempfile.mkdtemp()) / "nexonai-logo.png"
    im.save(out)
    return str(out)


def frame(slide, L, n, bg):
    rect(slide, 0, 0, W, H, fill=bg)
    # grey status strip
    rect(slide, 0, 0, W, 0.42, fill=SURFACE_MID)
    rect(slide, MX, 0.165, 0.09, 0.09, fill=PRIMARY, shape=MSO_SHAPE.OVAL)
    text(slide, MX + 0.18, 0.06, 4, 0.3, "unpruuf Compliance", size=10, color=INK, font=MONO, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 3.7, 0.06, 5.9, 0.3, L["status"], size=9.5, color=VARIANT, font=MONO, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, W - MX - 2, 0.06, 2, 0.3, L["lang"], size=10, color=PRIMARY, font=MONO, bold=True,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    # footer
    rect(slide, 0, 6.95, W, 0.55, fill=WHITE)
    rect(slide, 0, 6.95, W, 0.012, fill=LINE)
    pic = slide.shapes.add_picture(LOGO, Inches(MX), Inches(7.05), height=Inches(0.36))
    text(slide, MX + Emu(pic.width).inches + 0.2, 7.05, 4, 0.36, "NexonAI Consulting SRL", size=10, color=INK, font=SANS, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, W - MX - 5, 7.05, 5, 0.36, f"Alba Iulia  ·  unpruuf@nexonai.ro  ·  {n} / 6", size=9.5,
         color=VARIANT, font=MONO, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def heading(slide, lab, title, lead=None, y=0.72, title_size=30):
    label(slide, MX, y, 9, lab)
    text(slide, MX, y + 0.34, W - 2 * MX, 0.7, title, size=title_size, bold=True, color=INK)
    if lead:
        text(slide, MX, y + 1.02, W - 2 * MX, 0.4, lead, size=15, color=VARIANT)


# ---------------------------------------------------------------- slides
def s1(prs, L):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    frame(sl, L, 1, SURFACE)
    T = L["s1"]
    # navy hero block
    card(sl, MX, 0.72, W - 2 * MX, 3.85, fill=NAVY, line=None)
    label(sl, MX + 0.55, 1.12, 8, T["label"], color=SKY)
    text(sl, MX + 0.55, 1.5, 11, 1.0, "unpruuf Compliance", size=50, bold=True, color=WHITE)
    text(sl, MX + 0.55, 2.62, 11.2, 1.0, T["value"], size=18, color="dbe7ef", spacing=1.2)
    chip(sl, MX + 0.55, 3.88, T["chip"], fg=NAVY, bg=SKY, size=10.5, h=0.38)
    # three tiles
    tw = (W - 2 * MX - 2 * 0.3) / 3
    for i, (h, b) in enumerate(T["tiles"]):
        x = MX + i * (tw + 0.3)
        card(sl, x, 4.8, tw, 1.9)
        rect(sl, x + 0.35, 5.08, 0.42, 0.06, fill=PRIMARY)
        text(sl, x + 0.35, 5.25, tw - 0.7, 0.4, h, size=17, bold=True)
        text(sl, x + 0.35, 5.72, tw - 0.7, 0.85, b, size=13, color=VARIANT)
    notes(sl, T["notes"])


def phone(sl, x, y, w=1.15, h=2.05):
    rect(sl, x, y, w, h, fill=NAVY, radius=0.2)
    rect(sl, x + 0.09, y + 0.22, w - 0.18, h - 0.44, fill=WHITE, radius=0.06)
    rect(sl, x + w / 2 - 0.16, y + 0.09, 0.32, 0.05, fill="3d5566", radius=0.02)
    # a sealed envelope on the screen
    ex, ey, ew, eh = x + 0.24, y + 0.75, w - 0.48, 0.48
    rect(sl, ex, ey, ew, eh, fill=TINT, line=PRIMARY, line_w=1.25)
    arrow_v = sl.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(ex), Inches(ey), Inches(ew), Inches(eh * 0.55))
    arrow_v.rotation = 180
    arrow_v.fill.solid(); arrow_v.fill.fore_color.rgb = rgb(TINT)
    arrow_v.line.color.rgb = rgb(PRIMARY); arrow_v.line.width = Pt(1.25)


def laptop(sl, x, y, w=1.75, h=1.25):
    rect(sl, x + 0.12, y, w - 0.24, h, fill=NAVY, radius=0.08)
    rect(sl, x + 0.2, y + 0.08, w - 0.4, h - 0.16, fill=WHITE, radius=0.03)
    for i, c in enumerate((RED, AMBER, GREEN)):
        rect(sl, x + 0.3, y + 0.2 + i * 0.27, 0.12, 0.12, fill=c, shape=MSO_SHAPE.OVAL)
        rect(sl, x + 0.5, y + 0.23 + i * 0.27, w - 0.95, 0.06, fill=SURFACE_MID)
    rect(sl, x, y + h, w, 0.12, fill="3d5566", radius=0.05)


def server(sl, x, y, w, h, T):
    card(sl, x, y, w, h, fill=NAVY, line=None)
    text(sl, x, y + 0.2, w, 0.35, T["server"], size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(sl, x, y + 0.55, w, 0.3, T["server_sub"], size=11, color=SKY, align=PP_ALIGN.CENTER)
    # mailboxes, no numbers
    bw, bh, gx, gy = 0.5, 0.42, 0.16, 0.16
    cols, rows = 3, 2
    gw = cols * bw + (cols - 1) * gx
    sx, sy = x + (w - gw) / 2, y + 1.05
    for r in range(rows):
        for c in range(cols):
            bx, by = sx + c * (bw + gx), sy + r * (bh + gy)
            rect(sl, bx, by, bw, bh, fill="1d4258", line=SKY, radius=0.05, line_w=1)
            rect(sl, bx + 0.11, by + 0.12, bw - 0.22, 0.05, fill=SKY)
            if (r, c) == (0, 1):
                rect(sl, bx + 0.11, by + 0.2, bw - 0.22, 0.13, fill=WHITE)  # a sealed message waiting


def s2(prs, L):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    frame(sl, L, 2, SURFACE_LOW)
    T = L["s2"]
    heading(sl, T["label"], T["title"])
    # prerequisite banner
    rect(sl, MX, 1.82, W - 2 * MX, 0.52, fill=TINT, radius=0.1)
    rect(sl, MX, 1.82, 0.08, 0.52, fill=PRIMARY)
    text(sl, MX + 0.3, 1.82, W - 2 * MX - 0.5, 0.52,
         [[(T["pre_h"] + "  ", {"bold": True, "color": PRIMARY}), (T["pre"], {})]], size=14,
         anchor=MSO_ANCHOR.MIDDLE)
    # diagram card
    dx, dy, dw, dh = MX, 2.55, 7.35, 4.2
    card(sl, dx, dy, dw, dh)
    label(sl, dx + 0.3, dy + 0.18, 3, T["schema"], color=VARIANT, size=9)
    phone(sl, dx + 0.4, dy + 0.85)
    text(sl, dx + 0.05, dy + 3.0, 1.85, 0.6, [[(T["p1"], {"bold": True})], T["p1s"]], size=11.5,
         align=PP_ALIGN.CENTER)
    sx, sw = dx + 2.3, 2.55
    server(sl, sx, dy + 0.6, sw, 2.45, T)
    laptop(sl, dx + 5.62, dy + 1.25, w=1.5, h=1.1)
    text(sl, dx + 5.3, dy + 3.0, 2.0, 0.6, [[(T["p2"], {"bold": True})], T["p2s"]], size=11.5,
         align=PP_ALIGN.CENTER)
    ay = dy + 1.85
    arrow(sl, dx + 1.62, ay, sx - 0.06, ay)
    arrow(sl, dx + 5.62, ay, sx + sw + 0.06, ay)
    text(sl, dx + 1.5, ay - 0.62, 0.86, 0.55, T["a1"], size=10, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.BOTTOM)
    text(sl, sx + sw - 0.02, ay - 0.62, 0.9, 0.55, T["a2"], size=10, color=PRIMARY, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
    # NexonAI outside
    rect(sl, sx - 0.3, dy + 3.3, sw + 0.6, 0.55, fill=SURFACE_LOW, line=LINE, radius=0.1)
    text(sl, sx - 0.25, dy + 3.3, sw + 0.5, 0.55, T["nexon"], size=10.5, color=VARIANT, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    # points
    px, pw = dx + dw + 0.4, W - MX - (dx + dw + 0.4)
    y = 2.6
    for h, b in T["points"]:
        rect(sl, px, y + 0.08, 0.1, 0.1, fill=PRIMARY, shape=MSO_SHAPE.OVAL)
        text(sl, px + 0.25, y, pw - 0.25, 0.3, h, size=14, bold=True)
        text(sl, px + 0.25, y + 0.3, pw - 0.25, 0.55, b, size=11.5, color=VARIANT)
        y += 0.86
    # browser limit
    rect(sl, px, 6.03, pw, 0.72, fill=WARM_BG, radius=0.1)
    text(sl, px + 0.18, 6.03, pw - 0.3, 0.72, T["warn"], size=10, color=WARM, anchor=MSO_ANCHOR.MIDDLE)
    notes(sl, T["notes"])


def s3(prs, L):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    frame(sl, L, 3, SURFACE)
    T = L["s3"]
    heading(sl, T["label"], T["title"], T["lead"])
    n = len(T["steps"])
    gap = 0.22
    cw = (W - 2 * MX - (n - 1) * gap) / n
    y0 = 2.35
    for i, (h, deadline, b) in enumerate(T["steps"]):
        x = MX + i * (cw + gap)
        card(sl, x, y0, cw, 3.25)
        rect(sl, x + 0.28, y0 + 0.3, 0.5, 0.5, fill=PRIMARY if i not in (2, 4) else NAVY, shape=MSO_SHAPE.OVAL)
        text(sl, x + 0.28, y0 + 0.3, 0.5, 0.5, str(i + 1), size=16, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(sl, x + 0.28, y0 + 0.98, cw - 0.38, 0.65, h, size=13.5, bold=True, anchor=MSO_ANCHOR.TOP)
        if deadline:
            chip(sl, x + 0.28, y0 + 1.66, deadline, size=9.5, h=0.3, pad=0.12)
        text(sl, x + 0.28, y0 + (2.08 if deadline else 1.66), cw - 0.5, 1.0, b, size=11.5, color=VARIANT)
    # timeline
    ty = 5.95
    rect(sl, MX + 0.2, ty, W - 2 * MX - 0.4, 0.05, fill=LINE)
    marks = [(MX + 0.2, T["t"][0]), (MX + 2 * (cw + gap) + cw / 2, T["t"][1]),
             (MX + 4 * (cw + gap) + cw / 2, T["t"][2])]
    for mx_, t in marks:
        rect(sl, mx_ - 0.09, ty - 0.065, 0.18, 0.18, fill=NAVY, shape=MSO_SHAPE.OVAL)
        text(sl, mx_ - 1, ty + 0.18, 2, 0.3, t, size=10, color=NAVY, font=MONO, bold=True, align=PP_ALIGN.CENTER)
    text(sl, MX, 6.5, W - 2 * MX, 0.35, T["note"], size=12, color=VARIANT, align=PP_ALIGN.CENTER)
    notes(sl, T["notes"])


def s4(prs, L):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    frame(sl, L, 4, SURFACE_LOW)
    T = L["s4"]
    heading(sl, T["label"], T["title"])
    # illustration
    ix, iy, iw, ih = MX, 1.75, 6.35, 4.1
    card(sl, ix, iy, iw, ih)
    chip(sl, ix + 0.25, iy + 0.22, T["illu"], fg=WARM, bg=WARM_BG, size=9, h=0.3, pad=0.12)
    rect(sl, ix + 0.25, iy + 0.7, iw - 0.5, 0.45, fill=NAVY, radius=0.06)
    text(sl, ix + 0.45, iy + 0.7, 3, 0.45, T["mock_title"], size=12, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    cols = [ix + 0.45, ix + 2.2, ix + 3.75]
    for c, t in zip(cols, T["mock_cols"]):
        text(sl, c, iy + 1.3, 2, 0.25, t.upper(), size=8.5, color=VARIANT, font=MONO, bold=True)
    for r, (code, cat, dl, col) in enumerate(T["mock_rows"]):
        ry = iy + 1.62 + r * 0.56
        rect(sl, ix + 0.25, ry, iw - 0.5, 0.48, fill=SURFACE if r % 2 == 0 else WHITE, line=LINE, radius=0.06)
        text(sl, cols[0], ry, 1.8, 0.48, code, size=10, font=MONO, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(sl, cols[1], ry, 1.6, 0.48, cat, size=10.5, anchor=MSO_ANCHOR.MIDDLE)
        rect(sl, cols[2], ry + 0.18, 0.12, 0.12, fill=col, shape=MSO_SHAPE.OVAL)
        text(sl, cols[2] + 0.2, ry, 2.2, 0.48, dl, size=10.5, color=col, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    by = iy + 0.76
    for j, (t, fill, fg, ln) in enumerate(((T["btn_export"], WHITE, PRIMARY, PRIMARY),
                                          (T["btn_delete"], WHITE, RED, RED))):
        bx = ix + iw - 0.35 - (2 - j) * 1.15
        rect(sl, bx, by, 1.05, 0.33, fill=fill, line=ln, radius=0.08)
        text(sl, bx, by, 1.05, 0.33, t, size=10.5, bold=True, color=fg, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    # features
    fx = ix + iw + 0.45
    fw = W - MX - fx
    y = 1.75
    for h, b in T["feat"]:
        rect(sl, fx, y + 0.04, 0.06, 0.78, fill=PRIMARY)
        text(sl, fx + 0.25, y, fw - 0.25, 0.3, h, size=14.5, bold=True)
        text(sl, fx + 0.25, y + 0.32, fw - 0.25, 0.65, b, size=11.5, color=VARIANT)
        y += 1.04
    # GDPR line
    rect(sl, MX, 6.05, W - 2 * MX, 0.7, fill=WHITE, line=LINE, radius=0.1)
    text(sl, MX + 0.3, 6.05, W - 2 * MX - 0.6, 0.7, T["gdpr"], size=11, color=VARIANT, anchor=MSO_ANCHOR.MIDDLE)
    notes(sl, T["notes"])


def s5(prs, L):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    frame(sl, L, 5, SURFACE)
    T = L["s5"]
    heading(sl, T["label"], T["title"], T["lead"])
    gap = 0.3
    cw = (W - 2 * MX - 3 * gap) / 4
    for i, (lab, big, b) in enumerate(T["cards"]):
        x = MX + i * (cw + gap)
        card(sl, x, 2.35, cw, 3.3)
        label(sl, x + 0.3, 2.62, cw - 0.6, lab, size=10)
        text(sl, x + 0.3, 2.95, cw - 0.6, 0.8, big, size=20, bold=True, color=NAVY)
        text(sl, x + 0.3, 3.85, cw - 0.6, 1.7, b, size=12, color=VARIANT)
    text(sl, MX, 5.88, W - 2 * MX, 0.35, T["nocert"], size=13, color=INK, bold=True)
    text(sl, MX, 6.32, W - 2 * MX, 0.5, T["source"], size=9.5, color=VARIANT, font=MONO)
    notes(sl, T["notes"])


def s6(prs, L):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    frame(sl, L, 6, SURFACE_LOW)
    T = L["s6"]
    heading(sl, T["label"], T["title"], T["lead"])
    lx, lw = MX, 7.0
    card(sl, lx, 2.3, lw, 4.4)
    label(sl, lx + 0.35, 2.52, 6, T["inc"], size=10)
    y = 2.92
    for item in T["items"]:
        rect(sl, lx + 0.38, y + 0.1, 0.16, 0.16, fill=TINT, line=PRIMARY, radius=0.03, line_w=1)
        text(sl, lx + 0.7, y, lw - 1.0, 0.5, item, size=12.5)
        y += 0.52
    rx = lx + lw + 0.35
    rw = W - MX - rx
    card(sl, rx, 2.3, rw, 1.95, fill=NAVY, line=None)
    label(sl, rx + 0.35, 2.52, rw - 0.6, T["trial_l"], color=SKY, size=10)
    text(sl, rx + 0.35, 2.85, rw - 0.6, 0.5, T["trial_h"], size=24, bold=True, color=WHITE)
    text(sl, rx + 0.35, 3.42, rw - 0.6, 0.75, T["trial_b"], size=12, color="dbe7ef")
    card(sl, rx, 4.45, rw, 2.25)
    label(sl, rx + 0.35, 4.67, rw - 0.6, T["next_l"], size=10)
    text(sl, rx + 0.35, 5.0, rw - 0.6, 0.5, T["next_h"], size=19, bold=True)
    text(sl, rx + 0.35, 5.55, rw - 0.6, 0.35, T["next_b"], size=12, color=VARIANT)
    text(sl, rx + 0.35, 5.95, rw - 0.6, 0.45, "unpruuf@nexonai.ro", size=20, bold=True, color=PRIMARY)
    notes(sl, T["notes"])


# ---------------------------------------------------------------- texts
EN = {
    "lang": "EN",
    "status": "Tested · not independently audited · not certified",
    "s1": {
        "label": "Short presentation",
        "value": "Your own internal reporting channel for Romanian organisations with 50 or more employees: "
                 "on your own server, without collecting the identity of the person who reports.",
        "chip": "Tested · not independently audited",
        "tiles": [
            ("Your own server", "Your organisation runs the server. NexonAI runs no servers and has no access."),
            ("No reporter identities", "No account, no name, no phone number, no e-mail address."),
            ("Ready in one week", "With training for your case officers and an evidence pack for your DPO."),
        ],
        "notes": "unpruuf Compliance gives your organisation its own internal reporting channel, as Law 361/2022 "
                 "asks of employers with 50 or more employees. It runs on your own server and does not collect who "
                 "reports. It is tested and works; it is not independently audited and not certified, and an "
                 "independent security test is planned.",
    },
    "s2": {
        "label": "How unpruuf works",
        "title": "Two parts: an app and your own server.",
        "pre_h": "Prerequisite:",
        "pre": "your organisation runs its own server. NexonAI runs no servers and has no access.",
        "schema": "Schematic",
        "p1": "Employee", "p1s": "app or browser",
        "p2": "Case officer", "p2s": "case officer app",
        "server": "Your server", "server_sub": "my server, my rules",
        "a1": "leaves a sealed message", "a2": "picks it up",
        "nexon": "NexonAI: no server, no access, no copies",
        "points": [
            ("No identities", "No accounts, no phone numbers, no user directory."),
            ("Nobody writes directly", "You leave a sealed message on your own server. Your contact picks it up "
                                       "there. Replies take the same route."),
            ("My server, my rules", "The server belongs to your organisation. Nobody in the middle."),
            ("Nothing lies around", "What is not picked up is gone after 24 hours at most. A passage, not a "
                                    "storage."),
        ],
        "warn": "Reporting in the browser does not use Tor: your server sees the IP address. The first screen "
                "says so. The app does not have this limit.",
        "notes": "Everything rests on one condition: your organisation runs the server, NexonAI runs none and has "
                 "no access. Nobody writes directly to anyone: a sealed message is left on your server and picked "
                 "up there, and nothing stays longer than 24 hours. The browser option is weaker, because it shows "
                 "the IP address to your server; it says so on its first screen.",
    },
    "s3": {
        "label": "What your organisation gets",
        "title": "A report, from the first message to closing.",
        "lead": "Five steps. Your case officers always see the next deadline.",
        "steps": [
            ("Report", None, "An employee reports in the app or in the browser, from the QR code on the "
                             "workplace notice. No name needed."),
            ("Case code", None, "The reporter receives a random case code and follows the case with it."),
            ("Acknowledgement", "within 7 days", "The case officer confirms receipt of the report."),
            ("Dialogue", None, "The case officer asks questions. The reporter answers without giving contact "
                               "details."),
            ("Feedback, close", "within 3 months", "The case officer reports back on the measures taken and "
                                                   "closes the case."),
        ],
        "t": ["Day 0", "7 days", "3 months"],
        "note": "Reporting page and QR code in Romanian and English.",
        "notes": "This is what your people and your case officers do with it. An employee reports, receives a case "
                 "code and can follow the case with it. Your case officer confirms within 7 days, can ask questions, "
                 "and reports back within 3 months.",
    },
    "s4": {
        "label": "Case handling and deadlines",
        "title": "Every case in one overview.",
        "illu": "Illustration",
        "mock_title": "Cases",
        "mock_cols": ["Case code", "Category", "Next deadline"],
        "mock_rows": [
            ("HW-M8ZC-4NQ7", "Conflict of interest", "Acknowledge · today", RED),
            ("HW-7Q4M-2X9D", "Fraud", "Acknowledge · 2 days", AMBER),
            ("HW-K3PA-9T6R", "Safety", "Feedback · 74 days", GREEN),
            ("HW-D4XV-8KMC", "Harassment", "Feedback · 81 days", GREEN),
        ],
        "btn_export": "Export", "btn_delete": "Delete",
        "feat": [
            ("Deadlines at a glance", "Open cases sorted by the next deadline, marked red, amber or green."),
            ("Delete a case", "The case officer types the case code to confirm. A deletion log keeps only the "
                              "case code and the time, no content."),
            ("Export a case", "A printable file in Romanian or English with dates, deadlines, status changes and "
                              "the whole conversation."),
            ("Evidence pack for your DPO", "Technical overview, technical and organisational measures, impact "
                                           "assessment template, data processing agreement template."),
        ],
        "gdpr": "The architecture supports data minimisation, storage limitation and confidentiality "
                "(GDPR Art. 5, 25, 32). Your organisation is the controller; your DPO or lawyer confirms.",
        "notes": "The overview sorts open cases by the next deadline. A case can be deleted after typing its code "
                 "again, and the deletion log shows that and when, without content; copies made outside the app, "
                 "such as backups or exported files, follow your own retention rules. Your DPO receives an evidence "
                 "pack, and your DPO or lawyer confirms the GDPR assessment for your organisation.",
    },
    "s5": {
        "label": "Law 361/2022 in brief",
        "title": "What the law asks of your organisation.",
        "lead": "Romania transposed Directive (EU) 2019/1937 with Law 361/2022.",
        "cards": [
            ("Who", "50+ employees", "Private employers with 50 or more employees need an internal reporting "
                                     "channel. For 50–249 employees, this applies since 17.12.2023."),
            ("Deadlines", "7 days · 3 months", "Acknowledge receipt within 7 days. Give feedback to the reporter "
                                               "within 3 months."),
            ("Register", "5 years", "Reports are kept in a register for 5 years."),
            ("Authority", "ANI", "The National Integrity Agency is the external reporting channel and the "
                                 "sanctioning authority."),
        ],
        "nocert": "The law does not require a certification of the reporting channel.",
        "source": "Source: Legea 361/2022, art. 7, 9, 10, 36  ·  checked 30.10.2026  ·  not legal advice, "
                  "your lawyer confirms",
        "notes": "A short, factual summary of what Law 361/2022 asks. The internal channel is required for private "
                 "employers with 50 or more employees, for those with 50 to 249 since 17 December 2023. The channel "
                 "itself needs no certification; your lawyer confirms how the law applies to your organisation.",
    },
    "s6": {
        "label": "Start in one week",
        "title": "unpruuf Compliance Start",
        "lead": "Everything you need to run your own channel, ready in one week.",
        "inc": "The package includes",
        "items": [
            "One licence for your server",
            "Installation of the server and the case officer app, remotely or on site",
            "Reporting page and QR code in Romanian and English",
            "2-hour training for your case officers",
            "Templates: workplace notice and employee information",
            "Evidence pack for your DPO",
            "One year of updates and e-mail support",
        ],
        "trial_l": "Free trial",
        "trial_h": "30 days free",
        "trial_b": "Your case officers try reporting, dialogue and deadlines with test reports. Then you decide.",
        "next_l": "Next step",
        "next_h": "20-minute online demo",
        "next_b": "Write to us and pick a time.",
        "notes": "The package covers installation, training, templates and the evidence pack, and your channel is "
                 "ready in one week. You can try it free for 30 days with test reports before you decide. The next "
                 "step is a 20-minute online demo; write to unpruuf@nexonai.ro.",
    },
}

RO = {
    "lang": "RO",
    "status": "Testat · neauditat independent · necertificat",
    "s1": {
        "label": "Prezentare scurtă",
        "value": "Propriul canal intern de raportare pentru organizațiile din România cu 50 de angajați sau mai "
                 "mult: pe serverul tău, fără a colecta identitatea celui care raportează.",
        "chip": "Testat · neauditat independent",
        "tiles": [
            ("Serverul tău", "Organizația ta operează serverul. NexonAI nu operează servere și nu are acces."),
            ("Fără identități", "Fără cont, fără nume, fără număr de telefon, fără e-mail."),
            ("Gata într-o săptămână", "Cu instruire pentru responsabilii de caz și pachet de dovezi pentru DPO."),
        ],
        "notes": "unpruuf Compliance îi oferă organizației tale propriul canal intern de raportare, așa cum cere "
                 "Legea 361/2022 angajatorilor cu 50 de angajați sau mai mult. Rulează pe serverul tău și nu "
                 "colectează cine raportează. Este testat și funcționează; nu este auditat independent și nu este "
                 "certificat, iar un test de securitate independent este planificat.",
    },
    "s2": {
        "label": "Cum funcționează unpruuf",
        "title": "Două părți: o aplicație și propriul tău server.",
        "pre_h": "Condiție:",
        "pre": "organizația ta își operează propriul server. NexonAI nu operează servere și nu are acces.",
        "schema": "Schemă",
        "p1": "Angajat", "p1s": "aplicație sau browser",
        "p2": "Responsabil de caz", "p2s": "aplicația responsabilului",
        "server": "Serverul tău", "server_sub": "serverul meu, regulile mele",
        "a1": "lasă un mesaj sigilat", "a2": "îl ridică",
        "nexon": "NexonAI: fără server, fără acces, fără copii",
        "points": [
            ("Fără identități", "Fără conturi, fără număr de telefon, fără listă de utilizatori."),
            ("Nimeni nu scrie direct", "Lași mesajul sigilat pe propriul server. Contactul tău îl ridică de "
                                       "acolo. Răspunsul merge pe același drum."),
            ("Serverul meu, regulile mele", "Serverul aparține organizației tale. Nimeni din mijloc."),
            ("Nu rămâne nimic", "Ce nu este ridicat dispare după cel mult 24 de ore. Un loc de trecere, nu un "
                                "depozit."),
        ],
        "warn": "Raportarea din browser nu folosește Tor: serverul tău vede adresa IP. Primul ecran spune asta. "
                "Aplicația nu are această limită.",
        "notes": "Totul pornește de la o condiție: organizația ta operează serverul, NexonAI nu operează niciunul "
                 "și nu are acces. Nimeni nu scrie direct cuiva: un mesaj sigilat este lăsat pe serverul tău și "
                 "ridicat de acolo, iar nimic nu rămâne mai mult de 24 de ore. Varianta din browser este mai slabă, "
                 "pentru că arată adresa IP serverului tău; spune asta pe primul ecran.",
    },
    "s3": {
        "label": "Ce primește organizația ta",
        "title": "O raportare, de la primul mesaj la închidere.",
        "lead": "Cinci pași. Responsabilii de caz văd mereu următorul termen.",
        "steps": [
            ("Raportare", None, "Un angajat raportează în aplicație sau în browser, pornind de la codul QR de pe "
                                "afișul de la locul de muncă. Fără nume."),
            ("Cod de caz", None, "Sesizantul primește un cod de caz aleatoriu și urmărește cazul cu el."),
            ("Confirmare", "în 7 zile", "Responsabilul de caz confirmă primirea raportării."),
            ("Dialog", None, "Responsabilul pune întrebări. Sesizantul răspunde fără să-și dea datele de "
                             "contact."),
            ("Răspuns, închidere", "în 3 luni", "Responsabilul comunică măsurile luate și închide cazul."),
        ],
        "t": ["Ziua 0", "7 zile", "3 luni"],
        "note": "Pagina de raportare și codul QR, în română și engleză.",
        "notes": "Acesta este felul în care oamenii tăi și responsabilii de caz lucrează cu el. Un angajat "
                 "raportează, primește un cod de caz și poate urmări cazul cu el. Responsabilul confirmă în 7 zile, "
                 "poate pune întrebări și comunică răspunsul în 3 luni.",
    },
    "s4": {
        "label": "Gestionarea cazurilor și termene",
        "title": "Toate cazurile într-o singură privire.",
        "illu": "Ilustrație",
        "mock_title": "Cazuri",
        "mock_cols": ["Cod de caz", "Categorie", "Următorul termen"],
        "mock_rows": [
            ("HW-M8ZC-4NQ7", "Conflict de interese", "Confirmare · azi", RED),
            ("HW-7Q4M-2X9D", "Fraudă", "Confirmare · 2 zile", AMBER),
            ("HW-K3PA-9T6R", "Siguranță", "Răspuns · 74 de zile", GREEN),
            ("HW-D4XV-8KMC", "Hărțuire", "Răspuns · 81 de zile", GREEN),
        ],
        "btn_export": "Export", "btn_delete": "Șterge",
        "feat": [
            ("Termenele dintr-o privire", "Cazurile deschise, ordonate după următorul termen, marcate cu roșu, "
                                          "galben sau verde."),
            ("Ștergerea unui caz", "Responsabilul confirmă tastând codul cazului. Jurnalul de ștergeri păstrează "
                                   "doar codul și ora, fără conținut."),
            ("Exportul unui caz", "Un fișier tipăribil în română sau engleză, cu date, termene, schimbări de stare "
                                  "și întreaga conversație."),
            ("Pachet de dovezi pentru DPO", "Prezentare tehnică, măsuri tehnice și organizatorice, model de "
                                            "evaluare a impactului, model de acord de prelucrare."),
        ],
        "gdpr": "Arhitectura sprijină minimizarea datelor, limitarea stocării și confidențialitatea "
                "(GDPR art. 5, 25, 32). Organizația ta este operatorul; DPO-ul sau juristul tău confirmă.",
        "notes": "Prezentarea generală ordonează cazurile deschise după următorul termen. Un caz se poate șterge "
                 "după ce codul lui este tastat din nou, iar jurnalul de ștergeri arată că și când, fără conținut; "
                 "copiile făcute în afara aplicației, de exemplu copii de siguranță sau fișiere exportate, urmează "
                 "regulile tale de păstrare. DPO-ul primește un pachet de dovezi, iar DPO-ul sau juristul tău "
                 "confirmă evaluarea GDPR pentru organizația ta.",
    },
    "s5": {
        "label": "Legea 361/2022 pe scurt",
        "title": "Ce cere legea organizației tale.",
        "lead": "România a transpus Directiva (UE) 2019/1937 prin Legea 361/2022.",
        "cards": [
            ("Cine", "50+ angajați", "Angajatorii privați cu 50 de angajați sau mai mult au nevoie de un canal "
                                     "intern de raportare. Pentru 50–249 de angajați, din 17.12.2023."),
            ("Termene", "7 zile · 3 luni", "Confirmarea primirii în 7 zile. Informarea sesizantului în 3 luni."),
            ("Registru", "5 ani", "Raportările se păstrează 5 ani într-un registru."),
            ("Autoritate", "ANI", "Agenția Națională de Integritate este canalul extern de raportare și "
                                  "autoritatea care aplică sancțiunile."),
        ],
        "nocert": "Legea nu cere certificarea canalului de raportare.",
        "source": "Sursa: Legea 361/2022, art. 7, 9, 10, 36  ·  verificat la 30.10.2026  ·  nu este consultanță "
                  "juridică, juristul tău confirmă",
        "notes": "Un rezumat scurt și factual al cerințelor Legii 361/2022. Canalul intern este obligatoriu pentru "
                 "angajatorii privați cu 50 de angajați sau mai mult, iar pentru cei cu 50–249 de angajați din 17 "
                 "decembrie 2023. Canalul nu are nevoie de certificare; juristul tău confirmă cum se aplică legea "
                 "organizației tale.",
    },
    "s6": {
        "label": "Start într-o săptămână",
        "title": "unpruuf Compliance Start",
        "lead": "Tot ce îți trebuie pentru propriul canal, gata într-o săptămână.",
        "inc": "Pachetul conține",
        "items": [
            "O licență pentru serverul tău",
            "Instalarea serverului și a aplicației, la distanță sau la fața locului",
            "Pagina de raportare și codul QR, în română și engleză",
            "Două ore de instruire pentru responsabilii de caz",
            "Modele: afiș pentru locul de muncă și informare pentru angajați",
            "Pachetul de dovezi pentru DPO",
            "Un an de actualizări și suport prin e-mail",
        ],
        "trial_l": "Perioadă de probă",
        "trial_h": "30 de zile gratuit",
        "trial_b": "Responsabilii tăi încearcă raportarea, dialogul și termenele cu raportări de test. Apoi decizi.",
        "next_l": "Următorul pas",
        "next_h": "Demo online de 20 de minute",
        "next_b": "Scrie-ne și alege o oră.",
        "notes": "Pachetul cuprinde instalarea, instruirea, modelele și pachetul de dovezi, iar canalul tău este gata "
                 "într-o săptămână. Îl poți încerca gratuit 30 de zile cu raportări de test înainte să decizi. "
                 "Următorul pas este un demo online de 20 de minute; scrie-ne la unpruuf@nexonai.ro.",
    },
}

FORBIDDEN = ["safeline", "certified", "certificat", "gdpr-compliant", "gdpr-konform", "conform gdpr", "unbreakable",
             "first", "primul", "unique", "unic", "safer than", "mai sigur", "protects sources", "100 %", "100%",
             "anonymous", "anonim", "audited", "auditat", "€", "eur ", "price", "preț", "lei", "ron", "250",
             "protocol", "double ratchet", "4096"]
ALLOWED = ["not independently audited", "not certified", "neauditat independent", "necertificat",
           "nu este auditat independent", "nu este certificat", "certification", "certificarea", "primul ecran",
           "primul mesaj", "first screen", "first message"]


def check_words(prs, name):
    import re
    bad = []
    for i, sl in enumerate(prs.slides, 1):
        chunks = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
        chunks.append(sl.notes_slide.notes_text_frame.text)
        low = " ".join(chunks).lower()
        for a in ALLOWED:
            low = low.replace(a, " ")
        for w in FORBIDDEN:
            if re.search(r"(?<![a-zăâîșț])" + re.escape(w), low):
                bad.append((i, w))
    for i, w in bad:
        print(f"  FORBIDDEN in {name} slide {i}: {w!r}")
    return not bad


def build(L, out: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    for f in (s1, s2, s3, s4, s5, s6):
        f(prs, L)
    prs.core_properties.title = "unpruuf Compliance"
    prs.core_properties.author = "NexonAI Consulting SRL"
    ok = check_words(prs, out.name)
    prs.save(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out.parent), str(out)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print(("ok   " if ok else "WARN ") + str(out))
    return ok


LOGO = logo_png()

if __name__ == "__main__":
    ok = build(EN, HERE / "unpruuf_Compliance_Kurzpraesentation_EN.pptx")
    ok &= build(RO, HERE / "unpruuf_Compliance_Prezentare_Scurta_RO.pptx")
    sys.exit(0 if ok else 1)
