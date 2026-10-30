"""Build NexonAI-styled DOCX + PDF from the small Markdown sources in this folder.

Usage: python3 build_docs.py <source.md> [...] --out <dir>
Front matter (between --- lines): lang (en/ro/de), label, title, subtitle, date, version, status, conf.
Body: '# ' heading 1, '## ' heading 2, '- ' bullet, '1. ' numbered step, '> ' note box,
'| a | b |' table (first row = header), blank line = paragraph break, **bold** inline.
"""
from __future__ import annotations

import copy
import re
import subprocess
import sys
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
BLUE = "307098"
GRID = "D5DEE5"


def parse(path: Path) -> tuple[dict, list[str]]:
    text = path.read_text(encoding="utf-8")
    _, fm, body = text.split("---\n", 2)
    meta = {}
    for line in fm.splitlines():
        if ":" in line:
            k, v = line.split(":", 1)
            meta[k.strip()] = v.strip()
    return meta, body.splitlines()


def add_runs(par, text: str, bold_all: bool = False, color: str | None = None):
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = par.add_run(part)
        r.bold = bold_all or i % 2 == 1
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), GRID)
        b.append(e)
    tblPr.append(b)


def add_table(doc, rows: list[list[str]]):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.autofit = False
    borders(t)
    # column widths follow the average text length, so long text columns get the room
    weights = [max(6, len(rows[0][j]) + 2, sum(len(r[j]) for r in rows if j < len(r)) / len(rows)) for j in range(len(rows[0]))]
    widths = [Cm(16.5 * w / sum(weights)) for w in weights]
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips)))
    for j, w in enumerate(widths):
        for cell in t.columns[j].cells:
            cell.width = w
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            if i == 0:
                shade(cell, BLUE)
                add_runs(p, val, bold_all=True, color="FFFFFF")
            else:
                add_runs(p, val)
            for r in p.runs:
                r.font.size = Pt(9)
    doc.add_paragraph()


def build(src: Path, out_dir: Path) -> Path:
    meta, lines = parse(src)
    doc = docx.Document(str(HERE / "templates" / f"doc_{meta.get('lang', 'en')}.docx"))
    body = doc.element.body
    paras = doc.paragraphs
    paras[0].runs[0].text = meta.get("label", "DOCUMENT")
    for r in paras[0].runs[1:]:
        r.text = ""
    for p, key in ((paras[1], "title"), (paras[2], "subtitle")):
        p.runs[0].text = meta.get(key, "")
        for r in p.runs[1:]:
            r.text = ""
    info = doc.tables[0]
    for j, key in enumerate(("date", "version", "status", "conf")):
        cell = info.cell(1, j).paragraphs[0]
        cell.runs[0].text = meta.get(key, "")
        for r in cell.runs[1:]:
            r.text = ""
    # drop the template's sample content after the info table
    keep_until = info._tbl
    drop = False
    for el in list(body):
        if el is keep_until:
            drop = True
            continue
        if drop and el.tag != qn("w:sectPr"):
            body.remove(el)
    doc.add_paragraph()

    para: list[str] = []
    table: list[list[str]] = []

    def flush():
        if para:
            add_runs(doc.add_paragraph(style="Normal"), " ".join(para))
            para.clear()
        if table:
            add_table(doc, table)
            table.clear()

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("|"):
            if para:
                flush()
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", c) for c in cells):
                continue
            table.append(cells)
            continue
        if table:
            flush()
        if not line.strip():
            flush()
        elif line.startswith("## "):
            flush(); add_runs(doc.add_paragraph(style="Heading 2"), line[3:])
        elif line.startswith("# "):
            flush(); add_runs(doc.add_paragraph(style="Heading 1"), line[2:])
        elif line.startswith("- "):
            flush(); add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"\d+\. ", line):
            flush()
            num, rest = line.split(" ", 1)
            p = doc.add_paragraph(style="List Paragraph")
            add_runs(p, f"**{num}** {rest}")
        elif line.startswith("> "):
            flush(); add_runs(doc.add_paragraph(style="NX Note"), line[2:])
        else:
            para.append(line.strip())
    flush()
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / (src.stem + ".docx")
    doc.save(str(out))
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(out)],
                   check=True, capture_output=True, timeout=180)
    return out


if __name__ == "__main__":
    args = sys.argv[1:]
    out_dir = Path(args[args.index("--out") + 1])
    srcs = [Path(a) for a in args if a.endswith(".md")]
    for s in srcs:
        print(build(s, out_dir))
