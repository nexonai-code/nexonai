# -*- coding: utf-8 -*-
"""PNGs (render.js) -> PDF and PowerPoint (slides as images, speaker notes included).
Usage: python3 export2.py <png-root> <out-dir>   (png-root has one sub-folder per deck)"""
import glob, json, os, re, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

NAMES = {"business": "unpruuf_Business_RO", "business_de": "unpruuf_Business_DE",
         "compliance": "unpruuf_Compliance_RO", "compliance_de": "unpruuf_Compliance_DE"}
root, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
for deck, name in NAMES.items():
    files = sorted(glob.glob(os.path.join(root, deck, "*.png")))
    meta = json.load(open(os.path.join(root, deck, "notes.json"), encoding="utf-8"))
    imgs = [Image.open(f).convert("RGB") for f in files]
    imgs[0].save(os.path.join(out, name + ".pdf"), save_all=True, append_images=imgs[1:], resolution=144.0)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for f, sid in zip(files, meta["order"]):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(f, 0, 0, prs.slide_width, prs.slide_height)
        note = re.sub(r"<[^>]+>", "", meta["notes"].get(sid, "")).strip()
        if note:
            s.notes_slide.notes_text_frame.text = note
    prs.save(os.path.join(out, name + ".pptx"))
    print(name, len(files), "slides")
