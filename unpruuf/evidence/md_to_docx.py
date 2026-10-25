#!/usr/bin/env python3
"""Turns the audit briefs (plain Markdown: headings, paragraphs, bullets, numbered items, pipe tables,
code fences, **bold** and `code`) into .docx. Usage: python3 md_to_docx.py AUDIT_BRIEF_NODE_SERVER.md [more.md ...]
Needs python-docx. The .docx lands next to each .md."""
import re
import sys
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            par.add_run(f"{m.group(1)} ({m.group(2)})")
        else:
            par.add_run(part)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def convert(md_path):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    lines = open(md_path, encoding="utf-8").read().split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(block))
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(cols):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    par = cell.paragraphs[0]
                    add_inline(par, row[ci] if ci < len(row) else "")
                    for r in par.runs:
                        r.font.size = Pt(8.5)
                        if ri == 0:
                            r.bold = True
                    if ri == 0:
                        shade(cell, "E7ECEF")
            doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,3}) (.*)", line)
        if m:
            h = doc.add_heading(level=len(m.group(1)))
            add_inline(h, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\s*)- (.*)", line)
        if m:
            add_inline(doc.add_paragraph(style="List Bullet"), m.group(2))
            i += 1
            continue
        m = re.match(r"^\s*\d+\. (.*)", line)
        if m:
            add_inline(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue
        if line.strip():
            add_inline(doc.add_paragraph(), line.strip())
        i += 1
    out = re.sub(r"\.md$", ".docx", md_path)
    doc.save(out)
    return out


if __name__ == "__main__":
    for path in sys.argv[1:]:
        print(convert(path))
