# -*- coding: utf-8 -*-
"""A small drawing kit: one description of a slide, rendered twice, as an editable PPTX (python-pptx)
and as HTML (for the PDF). Units are inches on a 13.333 x 7.5 slide."""
import html as _h
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SW, SH = 13.333, 7.5
FONT = "Arial"; MONO = "Consolas"

def rgb(h): h = h.lstrip("#"); return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

class Slide:
    def __init__(self, deck, bg):
        self.deck = deck; self.bg = bg; self.items = []; self.notes = ""
        self.s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
        self.s.background.fill.solid(); self.s.background.fill.fore_color.rgb = rgb(bg)

    # ---- shapes ----
    def rect(self, x, y, w, h, fill=None, line=None, radius=0.0, lw=1.0, shadow=False, oval=False):
        shp = self.s.shapes.add_shape(MSO_SHAPE.OVAL if oval else (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE), Inches(x), Inches(y), Inches(w), Inches(h))
        if radius and not oval: shp.adjustments[0] = min(0.5, radius / min(w, h))
        if fill: shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
        else: shp.fill.background()
        if line: shp.line.color.rgb = rgb(line); shp.line.width = Pt(lw)
        else: shp.line.fill.background()
        shp.shadow.inherit = False
        if shadow:
            eff = shp._element.spPr.find(qn("a:effectLst"))
            if eff is None: eff = etree.SubElement(shp._element.spPr, qn("a:effectLst"))
            sh = etree.SubElement(eff, qn("a:outerShdw")); sh.set("blurRad", "228600"); sh.set("dist", "50800"); sh.set("dir", "5400000"); sh.set("algn", "t"); sh.set("rotWithShape", "0")
            c = etree.SubElement(sh, qn("a:srgbClr")); c.set("val", "0E2A2E"); a = etree.SubElement(c, qn("a:alpha")); a.set("val", "16000")
        self.items.append(("rect", x, y, w, h, fill, line, radius, lw, shadow, oval))
        return shp

    def text(self, x, y, w, h, paras, size=18, color="#0E2A2E", bold=False, font=FONT, align="l", anchor="t", space=0, shape=None):
        """paras: str | list of str | list of (str, dict)."""
        if isinstance(paras, str): paras = [paras]
        norm = [((p, {}) if isinstance(p, str) else p) for p in paras]
        box = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = box.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        for i, (t, o) in enumerate(norm):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[o.get("align", align)]
            sp = o.get("space", space)
            if sp: p.space_after = Pt(sp)
            r = p.add_run(); r.text = t; f = r.font
            f.name = o.get("font", font); f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold); f.color.rgb = rgb(o.get("color", color))
        self.items.append(("text", x, y, w, h, norm, dict(size=size, color=color, bold=bold, font=font, align=align, anchor=anchor, space=space)))
        return box

    def pic(self, path, x, y, w=None, h=None, shadow=False):
        from PIL import Image
        iw, ih = Image.open(path).size
        if w is None: w = h * iw / ih
        if h is None: h = w * ih / iw
        if shadow: self.rect(x + 0.03, y + 0.05, w - 0.06, h - 0.06, fill="#FFFFFF", radius=0.2, shadow=True)
        self.s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
        self.items.append(("pic", path, x, y, w, h))
        return w, h

    def arrow(self, x1, y1, x2, y2, color="#0F6B62", dashed=False, w=2.5):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color); c.line.width = Pt(w)
        ln = c.line._get_or_add_ln()
        if dashed: d = etree.SubElement(ln, qn("a:prstDash")); d.set("val", "dash")
        t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle")
        self.items.append(("arrow", x1, y1, x2, y2, color, dashed, w))

    def line(self, x1, y1, x2, y2, color="#D5DEDB", w=1.0):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color); c.line.width = Pt(w)
        self.items.append(("line", x1, y1, x2, y2, color, w))

class Deck:
    def __init__(self):
        self.prs = Presentation(); self.prs.slide_width = Emu(int(SW * 914400)); self.prs.slide_height = Emu(int(SH * 914400)); self.slides = []
    def slide(self, bg="#F2F5F4", notes=""):
        s = Slide(self, bg); s.notes = notes; self.slides.append(s); return s
    def save(self, path):
        for s in self.slides:
            if s.notes: s.s.notes_slide.notes_text_frame.text = s.notes
        self.prs.save(path)
    def html(self, path, base=""):
        px = lambda v: v * 96
        css = "@page{size:1280px 720px;margin:0}*{margin:0;box-sizing:border-box}body{background:#888}.sl{position:relative;width:1280px;height:720px;overflow:hidden;page-break-after:always;break-after:page}.a{position:absolute}.p{white-space:pre-wrap;overflow-wrap:anywhere}"
        parts = []
        for s in self.slides:
            h = []
            for it in s.items:
                k = it[0]
                if k == "rect":
                    _, x, y, w, hh, fill, line, radius, lw, shadow, oval = it
                    st = f"left:{px(x)}px;top:{px(y)}px;width:{px(w)}px;height:{px(hh)}px;"
                    if fill: st += f"background:{fill};"
                    if line: st += f"border:{lw*1.333:.1f}px solid {line};"
                    if oval: st += "border-radius:50%;"
                    elif radius: st += f"border-radius:{px(radius)}px;"
                    if shadow: st += "box-shadow:0 5px 22px rgba(14,42,46,.16);"
                    h.append(f'<div class="a" style="{st}"></div>')
                elif k == "pic":
                    _, pth, x, y, w, hh = it
                    h.append(f'<img class="a" src="{base}{pth}" style="left:{px(x)}px;top:{px(y)}px;width:{px(w)}px;height:{px(hh)}px">')
                elif k == "text":
                    _, x, y, w, hh, paras, d = it
                    anchor = {"t": "flex-start", "m": "center", "b": "flex-end"}[d["anchor"]]
                    ps = []
                    for t, o in paras:
                        size = o.get("size", d["size"]) * 1.333; col = o.get("color", d["color"]); b = o.get("bold", d["bold"]); fn = o.get("font", d["font"])
                        al = {"l": "left", "c": "center", "r": "right"}[o.get("align", d["align"])]; sp = o.get("space", d["space"]) or 0
                        ff = "Consolas,'DejaVu Sans Mono',monospace" if fn == MONO else "Arial,'Liberation Sans',sans-serif"
                        ps.append(f'<div class="p" style="font:{"700" if b else "400"} {size:.1f}px/1.22 {ff};color:{col};text-align:{al};margin-bottom:{sp*1.333:.0f}px">{_h.escape(t) or "&nbsp;"}</div>')
                    h.append(f'<div class="a" style="left:{px(x)}px;top:{px(y)}px;width:{px(w)}px;height:{px(hh)}px;display:flex;flex-direction:column;justify-content:{anchor};padding:3px 5px">{"".join(ps)}</div>')
                elif k == "arrow":
                    _, x1, y1, x2, y2, col, dashed, w = it
                    dash = 'stroke-dasharray="9 7"' if dashed else ""; mid = f"m{id(it)}"
                    h.append(f'<svg class="a" style="left:0;top:0" width="1280" height="720"><defs><marker id="{mid}" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="5" markerHeight="5" orient="auto"><path d="M0 0L10 5L0 10z" fill="{col}"/></marker></defs><line x1="{px(x1)}" y1="{px(y1)}" x2="{px(x2)}" y2="{px(y2)}" stroke="{col}" stroke-width="{w*1.333:.1f}" {dash} marker-end="url(#{mid})"/></svg>')
                elif k == "line":
                    _, x1, y1, x2, y2, col, w = it
                    h.append(f'<svg class="a" style="left:0;top:0" width="1280" height="720"><line x1="{px(x1)}" y1="{px(y1)}" x2="{px(x2)}" y2="{px(y2)}" stroke="{col}" stroke-width="{w*1.333:.1f}"/></svg>')
            parts.append(f'<div class="sl" style="background:{s.bg}">{"".join(h)}</div>')
        open(path, "w", encoding="utf-8").write(f'<!doctype html><meta charset="utf-8"><style>{css}</style><body>{"".join(parts)}</body>')
