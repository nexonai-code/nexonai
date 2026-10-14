# -*- coding: utf-8 -*-
import os, docx
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import content1, content2, content3, content4
C = content1.C
NAV=RGBColor(0x14,0x2B,0x4D)
def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr(); s=OxmlElement('w:shd')
    s.set(qn('w:val'),'clear'); s.set(qn('w:color'),'auto'); s.set(qn('w:fill'),hexc); tcPr.append(s)
def make(key,lang,d,out):
    D=docx.Document()
    sec=D.sections[0]; sec.left_margin=sec.right_margin=Cm(2.2); sec.top_margin=sec.bottom_margin=Cm(2)
    st=D.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
    p=D.add_paragraph(); r=p.add_run(d['title']); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=NAV
    p=D.add_paragraph(); r=p.add_run(d['sub']); r.italic=True; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0x55,0x55,0x55)
    for h,items in d['secs']:
        hp=D.add_paragraph(); hp.paragraph_format.space_before=Pt(12); hp.paragraph_format.keep_with_next=True
        r=hp.add_run(h); r.bold=True; r.font.size=Pt(13.5); r.font.color.rgb=NAV
        for it in items:
            if it[0]=="p":
                D.add_paragraph(it[1]).paragraph_format.space_after=Pt(6)
            elif it[0]=="b":
                for t in it[1]:
                    q=D.add_paragraph(t,style='List Bullet'); q.paragraph_format.space_after=Pt(2)
            else:
                hdr,rows=it[1],it[2]
                tb=D.add_table(rows=1,cols=len(hdr)); tb.style='Table Grid'
                for i,x in enumerate(hdr):
                    c=tb.rows[0].cells[i]; c.text=''; rr=c.paragraphs[0].add_run(x); rr.bold=True; rr.font.size=Pt(9.5); rr.font.color.rgb=RGBColor(255,255,255); shade(c,'142B4D')
                for row in rows:
                    cs=tb.add_row().cells
                    for i,x in enumerate(row):
                        cs[i].text=''; rr=cs[i].paragraphs[0].add_run(x); rr.font.size=Pt(9.5)
                        if i==0: rr.bold=True
                D.add_paragraph()
    D.save(out)
os.makedirs('out',exist_ok=True)
n=0
for key,langs in C.items():
    for l,d in langs.items():
        make(key,l,d,f"out/{key}_{l}.docx"); n+=1
print(n)
