import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
LOGO='/tmp/claude-0/-home-user-nexonai/dd345a09-6039-5c13-9173-147314f838d9/scratchpad/ci/NexonAI_CI_2026-10-19/01_Logo/Full_logo_transparent/nexonai-logo-800w.png'
BLUE=RGBColor(0x30,0x70,0x98);INK=RGBColor(0x17,0x27,0x33);GRAPH=RGBColor(0x50,0x5B,0x63)
SANS='Open Sans';MONO='DejaVu Sans Mono'
CONTACT='NexonAI Consulting SRL · Alba Iulia, Romania · gabriel@nexonai.ro · +40 773 829 646'
L={
'de':dict(lang='de-DE',doc='DOKUMENT',title='Dokumenttitel',sub='Untertitel oder Kurzbeschreibung',meta=['Datum','Version','Status','Vertraulichkeit'],mv=['1.0','Entwurf','Vertraulich'],
 h1='1. Zusammenfassung',p1='Hier steht der Einleitungstext. Er fasst in zwei bis vier Sätzen zusammen, worum es geht und was der Leser mitnehmen soll.',
 h2='2. Abschnitt',h3='2.1 Unterabschnitt',p2='Fließtext in Open Sans, 10,5 Punkt. Überschriften und Akzente sind im NexonAI-Blau gesetzt.',
 bl=['Erster Punkt der Liste','Zweiter Punkt der Liste','Dritter Punkt der Liste'],nl=['Erster Schritt','Zweiter Schritt','Dritter Schritt'],
 th=['Punkt','Beschreibung','Status'],tr=[['A','Kurze Beschreibung','Offen'],['B','Kurze Beschreibung','In Arbeit']],
 note='Hinweis',nt='Wichtige Anmerkung, Warnung oder Randbemerkung in einer hervorgehobenen Box.',page='Seite',of='von',
 lt='BRIEF',sender='NexonAI Consulting SRL · Alba Iulia · Romania',addr=['Empfänger / Firma','Name','Straße und Hausnummer','PLZ Ort, Land'],place='Alba Iulia',date='Datum',
 subj='Betreff',greet='Sehr geehrte Damen und Herren,',lbody='Hier steht der Brieftext. Kurze Absätze, klare Sprache, ein Anliegen pro Brief.',close='Mit freundlichen Grüßen',role='Gründer, NexonAI Consulting SRL'),
'en':dict(lang='en-GB',doc='DOCUMENT',title='Document title',sub='Subtitle or short description',meta=['Date','Version','Status','Confidentiality'],mv=['1.0','Draft','Confidential'],
 h1='1. Summary',p1='The introduction goes here. In two to four sentences it says what this is about and what the reader should take away.',
 h2='2. Section',h3='2.1 Subsection',p2='Body text in Open Sans, 10.5 pt. Headings and accents use the NexonAI blue.',
 bl=['First list item','Second list item','Third list item'],nl=['First step','Second step','Third step'],
 th=['Item','Description','Status'],tr=[['A','Short description','Open'],['B','Short description','In progress']],
 note='Note',nt='An important remark, warning or side note in a highlighted box.',page='Page',of='of',
 lt='LETTER',sender='NexonAI Consulting SRL · Alba Iulia · Romania',addr=['Recipient / Company','Name','Street and number','Postcode City, Country'],place='Alba Iulia',date='Date',
 subj='Subject',greet='Dear Sir or Madam,',lbody='The letter text goes here. Short paragraphs, clear language, one matter per letter.',close='Kind regards',role='Founder, NexonAI Consulting SRL'),
'ro':dict(lang='ro-RO',doc='DOCUMENT',title='Titlul documentului',sub='Subtitlu sau scurtă descriere',meta=['Data','Versiune','Stare','Confidențialitate'],mv=['1.0','Ciornă','Confidențial'],
 h1='1. Rezumat',p1='Aici se află textul introductiv. În două până la patru propoziții spune despre ce este vorba și ce trebuie să rețină cititorul.',
 h2='2. Secțiune',h3='2.1 Subsecțiune',p2='Text curent în Open Sans, 10,5 puncte. Titlurile și accentele folosesc albastrul NexonAI.',
 bl=['Primul punct din listă','Al doilea punct din listă','Al treilea punct din listă'],nl=['Primul pas','Al doilea pas','Al treilea pas'],
 th=['Punct','Descriere','Stare'],tr=[['A','Scurtă descriere','Deschis'],['B','Scurtă descriere','În lucru']],
 note='Notă',nt='O observație importantă, un avertisment sau o mențiune într-o casetă evidențiată.',page='Pagina',of='din',
 lt='SCRISOARE',sender='NexonAI Consulting SRL · Alba Iulia · România',addr=['Destinatar / Firmă','Nume','Strada și numărul','Cod poștal Localitate, Țara'],place='Alba Iulia',date='Data',
 subj='Subiect',greet='Stimată doamnă / Stimate domn,',lbody='Aici se află textul scrisorii. Paragrafe scurte, limbaj clar, o singură problemă per scrisoare.',close='Cu stimă,',role='Fondator, NexonAI Consulting SRL'),
}
def fonts(r,name,size=None,bold=None,color=None,italic=None):
    r.font.name=name
    rp=r._element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a),name)
    if size: r.font.size=Pt(size)
    if bold is not None: r.font.bold=bold
    if italic is not None: r.font.italic=italic
    if color is not None: r.font.color.rgb=color
def sfont(st,name,size,bold=False,color=INK,italic=False):
    st.font.name=name;st.font.size=Pt(size);st.font.bold=bold;st.font.italic=italic;st.font.color.rgb=color
    rp=st.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a),name)
    for a in ('w:asciiTheme','w:hAnsiTheme','w:eastAsiaTheme','w:cstheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
def spacing(st,before,after,line=1.35,keep=False):
    pf=st.paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after);pf.line_spacing=line
    if keep: pf.keep_with_next=True
def shade(el_pr,hexcol):
    s=OxmlElement('w:shd');s.set(qn('w:val'),'clear');s.set(qn('w:color'),'auto');s.set(qn('w:fill'),hexcol);el_pr.append(s)
def border(ppr,side,sz,color,space=4):
    b=ppr.find(qn('w:pBdr'))
    if b is None: b=OxmlElement('w:pBdr');ppr.append(b)
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),str(sz));e.set(qn('w:space'),str(space));e.set(qn('w:color'),color);b.append(e)
def field(par,instr,size=8,color=GRAPH):
    r=par.add_run();fonts(r,MONO,size,color=color)
    for t,txt in (('begin',None),('instr',instr),('sep',None),('text','1'),('end',None)):
        if t=='instr':
            e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=' '+instr+' ';r._element.append(e)
        elif t=='text':
            e=OxmlElement('w:t');e.text=txt;r._element.append(e)
        else:
            e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),{'begin':'begin','sep':'separate','end':'end'}[t]);r._element.append(e)
def setup(d,k):
    s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=s.right_margin=Cm(2.5);s.top_margin=Cm(3.2);s.bottom_margin=Cm(2.4);s.header_distance=Cm(1.3);s.footer_distance=Cm(1.2)
    s.different_first_page_header_footer=True
    st=d.styles
    sfont(st['Normal'],SANS,10.5);spacing(st['Normal'],0,6,1.35)
    sfont(st['Title'],SANS,26,True,INK);spacing(st['Title'],2,4,1.05)
    ppr=st['Title'].element.get_or_add_pPr()
    b=ppr.find(qn('w:pBdr'))
    if b is not None: ppr.remove(b)
    sfont(st['Subtitle'],SANS,13,False,GRAPH,False);spacing(st['Subtitle'],0,14,1.3)
    sfont(st['Heading 1'],SANS,18,True,BLUE);spacing(st['Heading 1'],14,6,1.15,True)
    sfont(st['Heading 2'],SANS,13.5,True,INK);spacing(st['Heading 2'],14,5,1.2,True)
    sfont(st['Heading 3'],SANS,11,True,BLUE);spacing(st['Heading 3'],10,3,1.2,True)
    for n in ('List Bullet','List Number'):
        sfont(st[n],SANS,10.5);spacing(st[n],0,3,1.35)
    from docx.enum.style import WD_STYLE_TYPE
    lab=st.add_style('NX Label',WD_STYLE_TYPE.PARAGRAPH);lab.base_style=st['Normal'];sfont(lab,MONO,8.5,False,BLUE);spacing(lab,0,6,1.1)
    lab.font.all_caps=True;lab.font.spacing=Pt(1)
    nb=st.add_style('NX Note',WD_STYLE_TYPE.PARAGRAPH);nb.base_style=st['Normal'];sfont(nb,SANS,10,False,INK);spacing(nb,8,8,1.4)
    nb.paragraph_format.left_indent=Cm(.4);nb.paragraph_format.right_indent=Cm(.2)
    p=nb.element.get_or_add_pPr();shade(p,'E3EEF6');border(p,'left',24,'307098',8)
    sm=st.add_style('NX Small',WD_STYLE_TYPE.PARAGRAPH);sm.base_style=st['Normal'];sfont(sm,SANS,8.5,False,GRAPH);spacing(sm,0,0,1.2)
    # language
    for sty in (st['Normal'],):
        rp=sty.element.get_or_add_rPr();lg=OxmlElement('w:lang');lg.set(qn('w:val'),L[k]['lang']);rp.append(lg)
def header_footer(d,k):
    s=d.sections[0];t=L[k]
    # first page header: large logo
    h=s.first_page_header;p=h.paragraphs[0];p.paragraph_format.space_after=Pt(0)
    p.add_run().add_picture(LOGO,width=Cm(4.6))
    # other header: small logo + right text
    h2=s.header;p2=h2.paragraphs[0];p2.paragraph_format.space_after=Pt(0)
    p2.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT)
    p2.add_run().add_picture(LOGO,width=Cm(2.6))
    r=p2.add_run('\tNexonAI Consulting SRL');fonts(r,MONO,8,color=GRAPH)
    border(p2._p.get_or_add_pPr(),'bottom',6,'D9DFE3',6)
    for f in (s.footer,s.first_page_footer):
        p=f.paragraphs[0];p.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT)
        border(p._p.get_or_add_pPr(),'top',6,'D9DFE3',6)
        r=p.add_run(CONTACT.replace('Romania','România') if k=='ro' else CONTACT);fonts(r,SANS,7.5,color=GRAPH)
        r=p.add_run('\t'+t['page']+' ');fonts(r,MONO,8,color=GRAPH);field(p,'PAGE')
        r=p.add_run(' '+t['of']+' ');fonts(r,MONO,8,color=GRAPH);field(p,'NUMPAGES')
def cell_fill(c,hexcol):
    tcPr=c._tc.get_or_add_tcPr();shade(tcPr,hexcol)
def cell_borders(c,color='D9DFE3'):
    tcPr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
    for side in ('top','left','bottom','right'):
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),color);b.append(e)
    tcPr.append(b)
def cell_text(c,text,bold=False,color=INK,size=9.5,mono=False):
    p=c.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
    r=p.add_run(text);fonts(r,MONO if mono else SANS,size,bold,color)
def margins(tbl):
    tblPr=tbl._tbl.tblPr;m=OxmlElement('w:tblCellMar')
    for side,w in (('top',60),('left',110),('bottom',60),('right',110)):
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(w));e.set(qn('w:type'),'dxa');m.append(e)
    tblPr.append(m)
def build_doc(k):
    t=L[k];d=Document();setup(d,k);header_footer(d,k)
    d.add_paragraph(t['doc'],style='NX Label')
    d.add_paragraph(t['title'],style='Title');d.add_paragraph(t['sub'],style='Subtitle')
    # meta table
    tb=d.add_table(rows=2,cols=4);tb.alignment=WD_TABLE_ALIGNMENT.LEFT;margins(tb)
    today=datetime.date(2026,10,19).strftime('%d.%m.%Y')
    vals=[today]+t['mv']
    for i in range(4):
        a,b=tb.cell(0,i),tb.cell(1,i)
        cell_fill(a,'EAF2F8');cell_fill(b,'EAF2F8')
        cell_text(a,t['meta'][i].upper(),False,BLUE,7.5,True);cell_text(b,vals[i],True,INK,10)
        for c in (a,b):
            cell_borders(c,'EAF2F8')
    d.add_paragraph().paragraph_format.space_after=Pt(2)
    d.add_paragraph(t['h1'],style='Heading 1');d.add_paragraph(t['p1'])
    d.add_paragraph(t['h2'],style='Heading 1');d.add_paragraph(t['h3'],style='Heading 2');d.add_paragraph(t['p2'])
    for x in t['bl']: d.add_paragraph(x,style='List Bullet')
    for x in t['nl']: d.add_paragraph(x,style='List Number')
    d.add_paragraph(t['h3'].replace('2.1','2.2'),style='Heading 2')
    tb=d.add_table(rows=1+len(t['tr']),cols=3);margins(tb)
    for i,hh in enumerate(t['th']):
        c=tb.cell(0,i);cell_fill(c,'307098');cell_borders(c,'307098');cell_text(c,hh,True,RGBColor(255,255,255),9.5)
    for ri,row in enumerate(t['tr'],1):
        for ci,v in enumerate(row):
            c=tb.cell(ri,ci);cell_borders(c);cell_text(c,v)
    for row in tb.rows:
        for c,w in zip(row.cells,(Cm(2.2),Cm(10.3),Cm(3.5))): c.width=w
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    n=d.add_paragraph(style='NX Note');r=n.add_run(t['note']+'  ');fonts(r,MONO,8.5,True,BLUE);n.add_run(t['nt'])
    d.core_properties.author='NexonAI Consulting SRL';d.core_properties.title=t['title']
    return d
def build_letter(k):
    t=L[k];d=Document();setup(d,k);header_footer(d,k)
    s=d.sections[0];s.top_margin=Cm(3.4)
    p=d.add_paragraph();r=p.add_run(t['sender']);fonts(r,SANS,7.5,color=GRAPH)
    p.paragraph_format.space_before=Pt(18);p.paragraph_format.space_after=Pt(4);border(p._p.get_or_add_pPr(),'bottom',4,'D9DFE3',2)
    for i,a in enumerate(t['addr']):
        q=d.add_paragraph();q.paragraph_format.space_after=Pt(0);rr=q.add_run(a);fonts(rr,SANS,10.5,bold=(i==0))
    q=d.add_paragraph();q.paragraph_format.space_before=Pt(18);q.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    today=datetime.date(2026,10,19).strftime('%d.%m.%Y')
    rr=q.add_run(f"{t['place']}, {today}");fonts(rr,SANS,10.5)
    q=d.add_paragraph();q.paragraph_format.space_before=Pt(14);rr=q.add_run(t['subj']);fonts(rr,SANS,12,True,BLUE)
    d.add_paragraph(t['greet']).paragraph_format.space_before=Pt(10)
    d.add_paragraph(t['lbody']);d.add_paragraph(t['lbody'])
    q=d.add_paragraph(t['close']);q.paragraph_format.space_before=Pt(12);q.paragraph_format.space_after=Pt(34)
    q=d.add_paragraph();q.paragraph_format.space_after=Pt(0);rr=q.add_run('Gabriel Hategan');fonts(rr,SANS,10.5,True)
    q=d.add_paragraph();rr=q.add_run(t['role']);fonts(rr,SANS,10,color=GRAPH)
    d.core_properties.author='NexonAI Consulting SRL';d.core_properties.title=t['subj']
    return d
import os
os.makedirs('out',exist_ok=True)
for k,name in (('de','Dokument'),('en','Document'),('ro','Document')):
    pass
names={'de':('Vorlage_Dokument_DE','Vorlage_Brief_DE'),'en':('Template_Document_EN','Template_Letter_EN'),'ro':('Sablon_Document_RO','Sablon_Scrisoare_RO')}
for k,(a,b) in names.items():
    build_doc(k).save(f'out/NexonAI_{a}.docx');build_letter(k).save(f'out/NexonAI_{b}.docx')
print(os.listdir('out'))
