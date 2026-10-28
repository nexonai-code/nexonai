#!/usr/bin/env python3
"""Turns our documents (plain Markdown: headings, paragraphs, bullets, numbered items, pipe tables,
code fences, **bold** and `code`) into .docx in the NexonAI / unpruuf design (design/NEXONAI_DESIGN.md).
Usage: python3 md_to_docx.py FILE.md [more.md ...]   Needs python-docx. The .docx lands next to each .md.
Open Sans is named as the font; where it is not installed Word shows its fallback, the layout stays."""
import re
import sys
from docx import Document
from docx.enum.text import WD_BREAK
import os
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = RGBColor(0x0F, 0x2A, 0x3A)
BLUE = RGBColor(0x30, 0x70, 0x98)
MUTED = RGBColor(0x4A, 0x5A, 0x66)
HERE = os.path.dirname(os.path.abspath(__file__))
SANS = "Open Sans"
MONO = "DejaVu Sans Mono"

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = MONO
            r.font.size = Pt(8.5)
            r.font.color.rgb = INK
            rpr = r._r.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F6F9")
            rpr.append(shd)
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            r = par.add_run(m.group(1))
            r.font.color.rgb = BLUE
            r.bold = True
            par.add_run(f" ({m.group(2)})").font.color.rgb = MUTED
        else:
            par.add_run(part)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_font(style, name, size, bold=None, color=None):
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def bottom_border(par, color="D9E4EA", size=6):
    p_pr = par._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(size)); b.set(qn("w:space"), "4"); b.set(qn("w:color"), color)
    borders.append(b)
    p_pr.append(borders)


def table_borders(table, color="D9E4EA"):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        borders.append(e)
    tbl_pr.append(borders)


def convert(md_path):
    doc = Document()
    set_font(doc.styles["Normal"], SANS, 10, color=INK)
    for name, size in (("Heading 1", 26), ("Heading 2", 17), ("Heading 3", 13)):
        set_font(doc.styles[name], SANS, size, bold=True, color=INK)
    for name in ("List Bullet", "List Number"):
        set_font(doc.styles[name], SANS, 10, color=INK)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    # Header: logo and the unpruuf wordmark, a hairline below, as on nexonai.ro.
    hp = sec.header.paragraphs[0]
    logo = os.path.join(HERE, "brand", "nexonai-logo.png")
    if os.path.exists(logo):
        hp.add_run().add_picture(logo, height=Cm(1.1))
    w1 = hp.add_run("   |   unpruuf")
    w1.font.name = SANS; w1.bold = True; w1.font.size = Pt(15); w1.font.color.rgb = INK
    w2 = hp.add_run(".")
    w2.font.name = SANS; w2.bold = True; w2.font.size = Pt(15); w2.font.color.rgb = BLUE
    bottom_border(hp)
    fp = sec.footer.paragraphs[0]
    # A file name ending in _RO means a Romanian document: same honesty line in Romanian.
    is_ro = os.path.splitext(os.path.basename(md_path))[0].endswith("_RO")
    fr = fp.add_run("NexonAI Consulting SRL, Alba Iulia  ·  fază pilot, neauditat independent" if is_ro
                    else "NexonAI Consulting SRL, Alba Iulia  ·  pilot phase, not independently audited")
    fr.font.name = SANS; fr.font.size = Pt(8); fr.font.color.rgb = MUTED
    lines = open(md_path, encoding="utf-8").read().split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(block))
            r.font.name = MONO
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0xE8, 0xF0, 0xF5)
            pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "0F2A3A")
            pr.append(shd)
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table_borders(table)
            for ri, row in enumerate(rows):
                for ci in range(cols):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    par = cell.paragraphs[0]
                    add_inline(par, row[ci] if ci < len(row) else "")
                    for r in par.runs:
                        r.font.size = Pt(8.5)
                        if r.font.name is None:
                            r.font.name = SANS
                        if ri == 0:
                            r.bold = True
                        elif ci == 0 and len(row[ci]) <= 14:
                            r.font.name = MONO
                            r.font.color.rgb = BLUE
                    if ri == 0:
                        shade(cell, "F2F6F9")
            doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,3}) (.*)", line)
        if m:
            h = doc.add_heading(level=len(m.group(1)))
            add_inline(h, m.group(2))
            for r in h.runs:
                r.font.name = SANS
                r.font.color.rgb = INK
            if len(m.group(1)) == 2:
                bottom_border(h)
            i += 1
            continue
        m = re.match(r"^(\s*)- (.*)", line)
        if m:
            add_inline(doc.add_paragraph(style="List Bullet"), m.group(2))
            i += 1
            continue
        m = re.match(r"^\s*\d+\. (.*)", line)
        if m:
            add_inline(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue
        if line.strip():
            add_inline(doc.add_paragraph(), line.strip())
        i += 1
    out = re.sub(r"\.md$", ".docx", md_path)
    doc.save(out)
    return out


if __name__ == "__main__":
    for path in sys.argv[1:]:
        print(convert(path))
