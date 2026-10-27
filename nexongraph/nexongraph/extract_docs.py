"""Document extractor: Markdown, text, reStructuredText, PDF, Word (.docx) and HTML.

Each document becomes a doc node with one section node per heading (or per PDF page).
Section text is kept for search and for the optional semantic pass.
"""
from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path

from .model import EXTRACTED, edge, new_extraction, node, ref, slug

MAX_SECTION_CHARS = 12000


def read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _sections_markdown(text: str) -> list[tuple[str, int, str, int]]:
    """Returns (title, level, body, line). Handles # headings and ===/--- underlined headings."""
    out: list[tuple[str, int, str, int]] = []
    cur_title, cur_level, cur_line, buf = None, 0, 1, []
    in_fence = False
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        title, level, skip = None, 0, 0
        if line.strip().startswith(("```", "~~~")):
            in_fence = not in_fence
        elif not in_fence:
            m = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
            if m:
                title, level = m.group(2).strip(), len(m.group(1))
            elif (i + 1 < len(lines) and line.strip() and not line.lstrip().startswith(("-", "*", "|", ">"))
                  and re.match(r"^(={3,}|-{3,})\s*$", lines[i + 1])):
                title, level, skip = line.strip(), (1 if lines[i + 1].startswith("=") else 2), 1
        if title:
            if cur_title is not None or "".join(buf).strip():
                out.append((cur_title or "", cur_level, "\n".join(buf).strip(), cur_line))
            cur_title, cur_level, cur_line, buf = title, level, i + 1, []
        else:
            buf.append(line)
        i += 1 + skip
    if cur_title is not None or "".join(buf).strip():
        out.append((cur_title or "", cur_level, "\n".join(buf).strip(), cur_line))
    return out


def _sections_rst(text: str) -> list[tuple[str, int, str, int]]:
    lines = text.splitlines()
    out, cur_title, cur_line, buf, levels = [], None, 1, [], {}
    i = 0
    while i < len(lines):
        line = lines[i]
        nxt = lines[i + 1] if i + 1 < len(lines) else ""
        if line.strip() and re.match(r"^([=\-~^\"'`#*+])\1{2,}\s*$", nxt) and len(nxt.strip()) >= len(line.strip()) - 2:
            ch = nxt.strip()[0]
            levels.setdefault(ch, len(levels) + 1)
            if cur_title is not None or "".join(buf).strip():
                out.append((cur_title or "", levels.get(ch, 1), "\n".join(buf).strip(), cur_line))
            cur_title, cur_line, buf = line.strip(), i + 1, []
            i += 2
            continue
        buf.append(line)
        i += 1
    if cur_title is not None or "".join(buf).strip():
        out.append((cur_title or "", 1, "\n".join(buf).strip(), cur_line))
    return out


def _sections_pdf(path: Path) -> tuple[list[tuple[str, int, str, int]], str | None]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return [], "pypdf missing"
    try:
        reader = PdfReader(str(path))
        title = None
        try:
            if reader.metadata and reader.metadata.title:
                title = str(reader.metadata.title)
        except Exception:
            pass
        out = []
        for i, page in enumerate(reader.pages, start=1):
            try:
                t = page.extract_text() or ""
            except Exception:
                t = ""
            t = t.strip()
            if not t:
                continue
            head = _guess_heading(t)
            out.append((f"S. {i}" + (f": {head}" if head else ""), 2, t, i))
        if not out:
            return [], "no text layer (scanned PDF?)"
        if title:
            out.insert(0, (title, 1, "", 1))
        return out, None
    except Exception as err:
        return [], f"pdf: {err}"


def _guess_heading(text: str) -> str | None:
    for line in text.splitlines()[:6]:
        s = line.strip()
        if 4 <= len(s) <= 80 and not s.endswith((".", ",", ";")) and re.search(r"[A-Za-zÄÖÜäöüăâîșțĂÂÎȘȚ]{3}", s):
            if re.match(r"^(\d+(\.\d+)*\.?\s+)?[A-ZÄÖÜĂÂÎȘȚ]", s):
                return s
    return None


def _sections_docx(path: Path) -> tuple[list[tuple[str, int, str, int]], str | None]:
    try:
        import docx  # python-docx
    except ImportError:
        return [], "python-docx missing"
    try:
        d = docx.Document(str(path))
    except Exception as err:
        return [], f"docx: {err}"
    out, cur_title, cur_level, buf, idx = [], None, 0, [], 1
    for i, p in enumerate(d.paragraphs, start=1):
        style = (p.style.name if p.style is not None else "") or ""
        text = p.text.strip()
        m = re.match(r"(Heading|Überschrift|Titlu|Title|Titel)\s*(\d*)", style, re.I)
        if m and text:
            if cur_title is not None or "".join(buf).strip():
                out.append((cur_title or "", cur_level, "\n".join(buf).strip(), idx))
            cur_title, cur_level, buf, idx = text, int(m.group(2) or 1), [], i
        elif text:
            buf.append(text)
    for table in d.tables:
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
        buf.append("\n".join(rows))
    if cur_title is not None or "".join(buf).strip():
        out.append((cur_title or "", cur_level, "\n".join(buf).strip(), idx))
    return out, None


class _HTMLSections(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out, self.title, self.level, self.buf, self.in_head = [], None, 0, [], None
        self.skip = 0
        self.links: list[str] = []
        self.doc_title = None
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript"):
            self.skip += 1
        elif tag in ("h1", "h2", "h3", "h4"):
            self.in_head = (int(tag[1]), [])
        elif tag == "title":
            self._in_title = True
        elif tag == "a":
            href = dict(attrs).get("href")
            if href:
                self.links.append(href)

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript") and self.skip:
            self.skip -= 1
        elif tag in ("h1", "h2", "h3", "h4") and self.in_head:
            lvl, parts = self.in_head
            title = " ".join("".join(parts).split())
            if self.title is not None or "".join(self.buf).strip():
                self.out.append((self.title or "", self.level, " ".join(self.buf).strip(), 0))
            self.title, self.level, self.buf, self.in_head = title, lvl, [], None
        elif tag == "title":
            self._in_title = False
        elif tag in ("p", "li", "div", "br", "tr"):
            self.buf.append("\n")

    def handle_data(self, data):
        if self.skip:
            return
        if self._in_title:
            self.doc_title = (self.doc_title or "") + data
        elif self.in_head:
            self.in_head[1].append(data)
        else:
            self.buf.append(data)

    def finish(self):
        if self.title is not None or "".join(self.buf).strip():
            self.out.append((self.title or "", self.level, re.sub(r"\n\s*\n+", "\n", " ".join(self.buf)).strip(), 0))
        return self.out


LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)\s]+)\)")


def extract_document(path: Path, rel: str, lang: str) -> dict:
    ex = new_extraction()
    error = None
    links: list[str] = []
    doc_title = None
    if lang == "markdown":
        text = read_text(path)
        sections = _sections_markdown(text)
        links = [m.group(2) for m in LINK_RE.finditer(text)]
    elif lang == "rst":
        sections = _sections_rst(read_text(path))
    elif lang == "text":
        text = read_text(path)
        sections = _sections_markdown(text) if re.search(r"^#{1,3}\s", text, re.M) else [("", 0, text.strip(), 1)]
    elif lang == "pdf":
        sections, error = _sections_pdf(path)
    elif lang == "docx":
        sections, error = _sections_docx(path)
    elif lang == "html":
        p = _HTMLSections()
        try:
            p.feed(read_text(path))
        except Exception as err:
            error = f"html: {err}"
        sections = p.finish()
        links = p.links
        doc_title = (p.doc_title or "").strip() or None
    else:
        sections = []

    first_heading = next((t for t, lvl, _, _ in sections if t and lvl <= 1), None)
    label = doc_title or first_heading or rel.split("/")[-1]
    did = node(ex, rel, label, "doc", rel, 1, lang=lang, file=rel.split("/")[-1], error=error)
    total_words = 0
    stack: list[tuple[int, str]] = []
    used: set[str] = set()
    for title, level, body, line in sections:
        body = body.strip()
        words = len(body.split())
        total_words += words
        if not title and not body:
            continue
        if not title:
            title = "Einleitung" if not stack else "(ohne Titel)"
        sid = f"{rel}#{slug(title)}"
        k = 2
        while sid in used:
            sid = f"{rel}#{slug(title)}-{k}"
            k += 1
        used.add(sid)
        node(ex, sid, title, "section", rel, line or None, lang=lang, level=level,
             text=body[:MAX_SECTION_CHARS], words=words)
        while stack and stack[-1][0] >= max(level, 1):
            stack.pop()
        parent = stack[-1][1] if stack else did
        edge(ex, parent, sid, "contains")
        stack.append((max(level, 1), sid))
    ex["nodes"][0]["words"] = total_words
    for href in links:
        if re.match(r"^[a-z]+://", href) or href.startswith(("#", "mailto:")):
            continue
        ref(ex, did, href.split("#")[0], "link")
    return ex
