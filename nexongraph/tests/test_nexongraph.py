import json
import textwrap
from pathlib import Path

import pytest

from nexongraph import config
from nexongraph.extract_docs import _sections_markdown, extract_document
from nexongraph.extract_python import extract_python
from nexongraph.extract_treesitter import extract_generic, extract_regex
from nexongraph.llm import parse_json
from nexongraph.mcp import Server
from nexongraph.pipeline import build
from nexongraph.query import Searcher, explain, load_graph, path, query_text


@pytest.fixture(autouse=True)
def home(tmp_path, monkeypatch):
    monkeypatch.setenv("NEXONGRAPH_HOME", str(tmp_path / "home"))


def make_project(root: Path) -> Path:
    (root / "app").mkdir(parents=True)
    (root / "web").mkdir()
    (root / "docs").mkdir()
    (root / "app" / "__init__.py").write_text("")
    (root / "app" / "store.py").write_text(textwrap.dedent('''
        """Case storage."""
        class CaseStore:
            def save_case(self, case):
                # WHY: cases are kept encrypted at rest
                return self.encrypt(case)

            def encrypt(self, data):
                return data
    '''))
    (root / "app" / "service.py").write_text(textwrap.dedent('''
        from app.store import CaseStore
        from . import store as st

        class Base:
            def run(self):
                pass

        class CaseService(Base):
            def __init__(self):
                self.store = CaseStore()

            def open_case(self, text):
                helper(text)
                return st.CaseStore().save_case(text)

        def helper(x):
            return x
    '''))
    (root / "web" / "client.ts").write_text(textwrap.dedent('''
        import { format } from "./util";
        export class RelayClient extends BaseClient {
          fetchMany(tags: string[]) { return format(tags); }
        }
    '''))
    (root / "web" / "util.ts").write_text("export function format(x: string[]) { return x.join(','); }\n")
    (root / "docs" / "guide.md").write_text(textwrap.dedent('''
        # Operations guide

        The `CaseService` opens cases and the RelayClient fetches packets.

        ## Storage

        Cases go through save_case() into the CaseStore. See [the client](../web/client.ts).
    '''))
    (root / "docs" / "notes.md").write_text("# Notes\n\nRelay packets and case storage are reviewed weekly. Relay packets expire.\n")
    return root


def test_python_extraction_nodes_and_refs():
    ex = extract_python("class A(B):\n    def m(self):\n        self.n()\n    def n(self):\n        pass\n", "x/a.py")
    ids = {n["id"] for n in ex["nodes"]}
    assert {"x/a.py", "x/a.py::A", "x/a.py::A.m", "x/a.py::A.n"} <= ids
    kinds = {(r["kind"], r["name"]) for r in ex["refs"]}
    assert ("inherit", "B") in kinds and ("call", "self.n") in kinds


def test_python_syntax_error_is_reported():
    ex = extract_python("def broken(:\n", "bad.py")
    assert "error" in ex["nodes"][0]


def test_typescript_extraction():
    src = "import { a } from './b';\nexport class C extends D { go() { a(); } }\nexport const f = () => g();\n"
    ex = extract_generic(src, "web/c.ts", "typescript")
    ids = {n["id"] for n in ex["nodes"]}
    assert "web/c.ts::C" in ids and "web/c.ts::C.go" in ids and "web/c.ts::f" in ids
    refs = {(r["kind"], r["name"]) for r in ex["refs"]}
    assert ("import", "./b") in refs and ("inherit", "D") in refs and ("call", "a") in refs


def test_regex_fallback():
    ex = extract_regex("class Foo {\n  fun bar() { baz() }\n}\nimport x.y.Z\n", "a.kt", "kotlin")
    labels = {n["label"] for n in ex["nodes"]}
    assert "Foo" in labels and "bar()" in labels


def test_markdown_sections():
    secs = _sections_markdown("intro\n# A\none\n## B\ntwo\n```\n# not a heading\n```\nC\n===\nthree")
    assert [s[0] for s in secs] == ["", "A", "B", "C"]
    assert "# not a heading" in secs[2][2]


def test_docx(tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("Handbook", 1)
    d.add_paragraph("Body text about relays.")
    p = tmp_path / "h.docx"
    d.save(p)
    ex = extract_document(p, "h.docx", "docx")
    assert any(n["kind"] == "section" and n["label"] == "Handbook" for n in ex["nodes"])


def test_parse_json_with_fences():
    assert parse_json('Sure:\n```json\n{"a": {"b": "}"}}\n```') == {"a": {"b": "}"}}


def test_full_build_and_queries(tmp_path):
    root = make_project(tmp_path / "proj")
    r = build(root, semantic=False)
    out = Path(r["out"])
    assert (out / "graph.json").exists() and (out / "graph.html").exists() and (out / "REPORT.md").exists()
    G = load_graph(out / "graph.json")
    e = lambda u, v: G.edges[u, v] if G.has_edge(u, v) else None
    # imports between files
    assert e("app/service.py", "app/store.py")["relation"] == "imports"
    assert e("web/client.ts", "web/util.ts")["relation"] == "imports"
    # calls: same file, imported symbol, module alias, self.method
    assert e("app/service.py::CaseService.open_case", "app/service.py::helper")["confidence"] == "EXTRACTED"
    assert e("app/service.py::CaseService.__init__", "app/store.py::CaseStore")["confidence"] == "EXTRACTED"
    assert e("app/service.py::CaseService.open_case", "app/store.py::CaseStore.save_case") is not None
    assert e("app/store.py::CaseStore.save_case", "app/store.py::CaseStore.encrypt")["confidence"] == "INFERRED"
    assert e("web/client.ts::RelayClient.fetchMany", "web/util.ts::format") is not None
    # inheritance inside the project, unknown base stays an attribute
    assert e("app/service.py::CaseService", "app/service.py::Base")["relation"] == "inherits"
    assert "BaseClient" in G.nodes["web/client.ts::RelayClient"].get("bases", [])
    # docs mention code and link to files
    assert e("docs/guide.md#operations-guide", "app/service.py::CaseService")["relation"] == "mentions"
    assert e("docs/guide.md", "web/client.ts")["relation"] == "links_to"
    # notes in code
    assert any(d.get("kind") == "note" and "encrypted" in d["label"] for _, d in G.nodes(data=True))
    # shared terms connect the two documents
    assert any(d.get("kind") == "concept" for _, d in G.nodes(data=True))
    # queries
    assert Searcher(G).best("CaseStore") == "app/store.py::CaseStore"
    assert explain(G, "RelayClient")["id"] == "web/client.ts::RelayClient"
    p = path(G, "Operations guide", "format")
    assert p and p["path"][0].startswith("docs/guide.md") and p["path"][-1] == "web/util.ts::format"
    assert "CaseStore" in query_text(G, "where are cases stored")
    # report and project registry
    assert "NexonGraph" in (out / "REPORT.md").read_text(encoding="utf-8")
    assert config.load()["projects"][0]["root"] == str(root.resolve())


def test_incremental_build_uses_cache(tmp_path):
    root = make_project(tmp_path / "proj")
    first = build(root, semantic=False)
    second = build(root, semantic=False)
    assert second["reused"] == first["files"]
    (root / "app" / "store.py").write_text("class CaseStore:\n    pass\n")
    third = build(root, semantic=False)
    assert third["reused"] == first["files"] - 1


def test_mcp_server(tmp_path):
    root = make_project(tmp_path / "proj")
    build(root, semantic=False)
    srv = Server(str(root))
    init = srv.handle({"jsonrpc": "2.0", "id": 1, "method": "initialize", "params": {}})
    assert init["result"]["serverInfo"]["name"] == "nexongraph"
    tools = srv.handle({"jsonrpc": "2.0", "id": 2, "method": "tools/list"})["result"]["tools"]
    assert {t["name"] for t in tools} >= {"query", "search", "explain", "path", "report"}
    res = srv.handle({"jsonrpc": "2.0", "id": 3, "method": "tools/call", "params": {"name": "explain", "arguments": {"name": "CaseStore"}}})
    assert "CaseStore" in res["result"]["content"][0]["text"]
    assert srv.handle({"jsonrpc": "2.0", "method": "notifications/initialized"}) is None


def test_semantic_pass_with_fake_model(tmp_path, monkeypatch):
    root = make_project(tmp_path / "proj")
    from nexongraph import llm

    def fake(cfg, system, user, json_mode=False, timeout=180, max_tokens=2000):
        return json.dumps({"entities": [{"name": "Relay packets", "type": "concept", "description": "d"},
                                        {"name": "Invented thing", "type": "concept"}],
                           "relations": [{"source": "Relay packets", "target": "Invented thing", "relation": "uses"}]})
    monkeypatch.setattr(llm, "complete", fake)
    cfg = config.load()
    cfg["llm"] = {"provider": "ollama", "model": "m", "api_key": "", "base_url": ""}
    r = build(root, semantic=True, cfg=cfg)
    G = load_graph(Path(r["out"]) / "graph.json")
    assert "concept::relay-packets" in G
    assert "concept::invented-thing" not in G  # not in the text: dropped
    assert r["semantic"].startswith("ollama")


def test_out_dir_in_profile(tmp_path):
    root = make_project(tmp_path / "proj")
    cfg = config.load()
    cfg["out_mode"] = "appdata"
    r = build(root, semantic=False, cfg=cfg)
    assert not (root / "nexongraph-out").exists()
    assert config.find_out_dir(root) == Path(r["out"])
